#!/usr/bin/env python3
"""Build output/g20-ministers-trade-reference.xlsx from data/g20.json."""
import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter

D = json.load(open('/home/user/Photo/data/g20.json'))
CO = D['countries']; meta = D['meta']
gy = meta['gdp']['reference_year']; ty = meta['us_trade']['reference_year']

NAVY = '1F3864'; NAVYD = '162B4D'
HDR_FILL = PatternFill('solid', fgColor=NAVY)
# Conditional-formatting differential fills render from bgColor — set both slots so the color shows.
RED = PatternFill(start_color='F8CBAD', end_color='F8CBAD', fill_type='solid'); RED_FT = Font(color='843C0C')
GRN = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid'); GRN_FT = Font(color='006100')
FLAG_FILL = {'Verified': PatternFill('solid', fgColor='E2EFDA'),
             'Needs check': PatternFill('solid', fgColor='FFF2CC'),
             'Vacant / acting': PatternFill('solid', fgColor='FCE4D6')}
STRIPE = PatternFill('solid', fgColor='F2F5FA')
thin = Side(style='thin', color='D9D9D9')
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)

wb = Workbook(); ws = wb.active; ws.title = 'G20 Ministers & Trade'

COLS = [
    ('Country', 18, 'text'),
    (f'GDP\n(USD bn, nominal, {gy})', 15, 'money'),
    ('GDP ref.\nyear', 9, 'center'),
    ('Overall trade balance\n(USD bn; goods basis*)', 17, 'money'),
    (f'U.S. goods exports\nto country (USD bn, {ty})', 15, 'money'),
    (f'U.S. goods imports\nfrom country (USD bn, {ty})', 15, 'money'),
    ('U.S. bilateral balance\n(USD bn; − = U.S. deficit)', 17, 'money'),
    ('Trade data\nref. year', 10, 'center'),
    ('Digital / technology\nminister', 22, 'text'),
    ('Digital minister title', 32, 'wrap'),
    ('Digital minister ministry', 30, 'wrap'),
    ('Trade / commerce\nminister', 22, 'text'),
    ('Trade minister title', 32, 'wrap'),
    ('Trade minister ministry', 30, 'wrap'),
    ('Confidence\nflag', 14, 'center'),
    ('Source notes', 66, 'wrap'),
]

# header row
for j,(title,width,_) in enumerate(COLS, start=1):
    c = ws.cell(row=1, column=j, value=title)
    c.font = Font(bold=True, color='FFFFFF', size=10)
    c.fill = HDR_FILL
    c.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    c.border = BORDER
    ws.column_dimensions[get_column_letter(j)].width = width
ws.row_dimensions[1].height = 42

def primary(offs):
    for o in offs:
        if o['is_primary']: return o
    return offs[0] if offs else {'name':'','title':'','ministry':''}

MONEY = '$#,##0.0;($#,##0.0)'
for i, cn in enumerate(CO):
    r = i + 2
    dg = primary(cn['digital_ministers']); tr = primary(cn['trade_ministers'])
    # append acting marker to name where acting
    dname = dg['name'] + (' (acting)' if dg.get('is_acting') else '')
    tname = tr['name'] + (' (acting)' if tr.get('is_acting') else '')
    vals = [
        cn['country'], cn['gdp_usd_b'], cn['gdp_year'], cn['overall_balance_usd_b'],
        cn['us_exports_usd_b'] if cn['us_exports_usd_b'] is not None else 'N/A',
        cn['us_imports_usd_b'] if cn['us_imports_usd_b'] is not None else 'N/A',
        cn['us_bilateral_usd_b'] if cn['us_bilateral_usd_b'] is not None else 'N/A',
        cn['us_trade_year'],
        dname, dg['title'], dg['ministry'],
        tname, tr['title'], tr['ministry'],
        cn['row_confidence'], cn['source_notes'],
    ]
    for j,(title,width,kind) in enumerate(COLS, start=1):
        c = ws.cell(row=r, column=j, value=vals[j-1])
        c.border = BORDER
        if kind == 'money' and isinstance(vals[j-1], (int,float)):
            c.number_format = MONEY
            c.alignment = Alignment(horizontal='right', vertical='center')
        elif kind == 'center':
            c.alignment = Alignment(horizontal='center', vertical='center')
        elif kind == 'wrap':
            c.alignment = Alignment(horizontal='left', vertical='top', wrap_text=True)
        else:
            c.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
        if i % 2 == 1 and j not in (7,15):
            c.fill = STRIPE
    # confidence flag fill
    ws.cell(row=r, column=15).fill = FLAG_FILL.get(cn['row_confidence'])
    ws.cell(row=r, column=15).font = Font(bold=True, size=9)
    ws.cell(row=r, column=1).font = Font(bold=True)
    ws.row_dimensions[r].height = 58

# conditional formatting: US bilateral balance (col G = 7)
last = len(CO) + 1
rng = f'G2:G{last}'
ws.conditional_formatting.add(rng, CellIsRule(operator='lessThan', formula=['0'], fill=RED, font=RED_FT))
ws.conditional_formatting.add(rng, CellIsRule(operator='greaterThan', formula=['0'], fill=GRN, font=GRN_FT))

# freeze header, autofilter, footnote
ws.freeze_panes = 'A2'
ws.auto_filter.ref = f'A1:{get_column_letter(len(COLS))}{last}'
fr = last + 2
ws.cell(row=fr, column=1,
    value=("*Overall trade balance is goods-only except Australia (goods+services) and India (FY2025-26). "
           "GDP: %s. U.S. bilateral goods trade: %s. Compiled %s. "
           "Flags: Verified / Needs check / Vacant-acting. Split portfolios: primary official shown; others in Source notes."
           % (meta['gdp']['source'].split('.')[0], meta['us_trade']['source'].split('.')[0], meta['compiled']))
    ).font = Font(italic=True, size=8, color='595959')
ws.merge_cells(start_row=fr, start_column=1, end_row=fr, end_column=16)
ws.cell(row=fr, column=1).alignment = Alignment(wrap_text=True, vertical='top')
ws.row_dimensions[fr].height = 44

wb.save('/home/user/Photo/output/g20-ministers-trade-reference.xlsx')
print("Wrote output/g20-ministers-trade-reference.xlsx —", len(CO), "rows")
