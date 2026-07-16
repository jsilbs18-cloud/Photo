#!/usr/bin/env python3
"""Build output/g20-ministers-trade-list.docx from data/g20.json — one page per
country, ITA Visual Style Guide format. Same source of truth as the deck."""
import json
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = '/home/user/Photo/'
D = json.load(open(ROOT + 'data/g20.json'))
CO = D['countries']; BL = D.get('blocs', [])

NAVY = RGBColor(0x0A, 0x31, 0x4D); BLUE = RGBColor(0x00, 0x55, 0x8C)
GOLD = RGBColor(0xB4, 0x86, 0x2D); INK = RGBColor(0x2C, 0x2C, 0x2C)
GRAY = RGBColor(0x54, 0x60, 0x6C); MUTED = RGBColor(0x75, 0x7D, 0x87)
ORANGE = RGBColor(0xD4, 0x61, 0x27); TEAL = RGBColor(0x00, 0x75, 0x82)
WHITE = RGBColor(0xFF, 0xFF, 0xFF); LTBLUE = RGBColor(0xB9, 0xC9, 0xD9)
SLATE = RGBColor(0xB1, 0xBB, 0xCA)
FONT = 'Open Sans'

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.left_margin = sec.right_margin = Inches(0.85)
sec.top_margin, sec.bottom_margin = Inches(0.7), Inches(0.8)

st = doc.styles['Normal']
st.font.name = FONT; st.font.size = Pt(10); st.font.color.rgb = INK
st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
st.paragraph_format.space_after = Pt(3)

def run(p, text, size=10, color=INK, bold=False, italic=False, caps_spacing=False):
    r = p.add_run(text); r.font.name = FONT; r.font.size = Pt(size)
    r.font.color.rgb = color; r.bold = bold; r.italic = italic
    r.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    if caps_spacing:
        sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), '24')
        r.element.rPr.append(sp)
    return r

def shade(p, hexcolor):
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexcolor)
    pPr.append(shd)

def bottom_border(p, hexcolor, size=12, space=4):
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement('w:pBdr'); bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single'); bot.set(qn('w:sz'), str(size))
    bot.set(qn('w:space'), str(space)); bot.set(qn('w:color'), hexcolor)
    pbdr.append(bot); pPr.append(pbdr)

def page_break_before(p):
    p._p.get_or_add_pPr().append(OxmlElement('w:pageBreakBefore'))

def tight(p, before=0, after=0):
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p

def _no_borders(tbl):
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement(f'w:{edge}'); e.set(qn('w:val'), 'none')
        e.set(qn('w:sz'), '0'); e.set(qn('w:space'), '0'); borders.append(e)
    tbl._tbl.tblPr.append(borders)

def money(v, year=None):
    """Prose form: spelled-out billions."""
    if v is None: return 'N/A'
    s = f"${abs(v):,.1f} billion"
    return s if year is None else f"{s} ({year})"

def money_b(v, year=None):
    """Standardized tabular form (matches the deck and workbook): $X,XXX.XB."""
    if v is None: return 'N/A'
    s = f"${abs(v):,.1f}B"
    return s if year is None else f"{s} ({year})"

def balance_word(v): return 'surplus' if (v or 0) >= 0 else 'deficit'

# ============================ COVER PAGE ============================
def navy_line(text=None, size=10, color=WHITE, bold=False, caps=False):
    p = doc.add_paragraph(); shade(p, '0A314D'); tight(p, 0, 0)
    p.paragraph_format.line_spacing = 1.0
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run(p, text if text is not None else ' ', size, color, bold=bold, caps_spacing=caps)
    return p

navy_line(size=16)
navy_line(size=10)
p = doc.add_paragraph(); shade(p, '0A314D'); tight(p, 0, 0)
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run().add_picture(ROOT + 'assets/seal/doc-seal-white.png', height=Inches(1.05))
navy_line(size=10)
navy_line('U.S. DEPARTMENT OF COMMERCE', 12, WHITE, bold=True, caps=True)
navy_line(size=2)
navy_line('INTERNATIONAL TRADE ADMINISTRATION', 12, LTBLUE, caps=True)
navy_line(size=22)
navy_line('G20 Digital & Trade Ministers', 30, WHITE, bold=True)
navy_line(size=6)
navy_line('GDP and U.S. Bilateral Trade Profiles', 16, SLATE)
navy_line(size=10)
navy_line('▬▬▬▬▬▬', 12, GOLD, bold=True)
navy_line(size=10)
navy_line('Internal Reference   |   Trade data: 2025 full year (goods basis)   |   July 16, 2026', 11, WHITE)
navy_line(size=16)
navy_line('One page per member: headline economic and U.S. trade figures, then the trade and', 9.5, LTBLUE)
navy_line('digital/technology officeholders — verified as current on 2026-07-16.', 9.5, LTBLUE)
navy_line(size=16)
navy_line(size=16)

ps = doc.add_paragraph(); tight(ps, 14, 2); run(ps, 'Scope', 13, NAVY, bold=True)
ps2 = doc.add_paragraph()
run(ps2, 'This reference profiles the 19 G20 sovereign member countries — with Poland covered in place of '
         'South Africa, a deliberate substitution — followed by an annex on the two G20 bloc members, the '
         'European Union and the African Union. Figures match the companion briefing deck; both documents '
         'are generated from the same verified dataset. Data caveats and the full uncertainty register are '
         'maintained separately (GAPS.md).', 10, INK)
ps3 = doc.add_paragraph()
run(ps3, 'Contents: one page per country (alphabetical) · EU & African Union annex.', 8.5, MUTED, italic=True)
pv = doc.add_paragraph()
run(pv, 'Sources: GDP — IMF WEO 2025 · U.S. goods trade — U.S. Census 2025 · officeholders verified 2026-07-16.',
    8.5, MUTED, italic=True)

# ============================ PROFILE WRITERS ============================
def country_heading(name, tag=None, flag=None, new_page=True):
    p = doc.add_paragraph(); tight(p, 0, 2)
    if new_page: page_break_before(p)
    bottom_border(p, 'B4862D', size=14, space=3)
    if flag:
        p.add_run().add_picture(ROOT + flag, height=Inches(0.24))
        run(p, '   ', 14)
    run(p, name, 14, NAVY, bold=True)
    if tag: run(p, f'    ·    {tag}', 9, TEAL, bold=True, caps_spacing=True)
    p.paragraph_format.keep_with_next = True

def subhead(text):
    p = doc.add_paragraph(); tight(p, 6, 1)
    p.paragraph_format.keep_with_next = True
    run(p, text.upper(), 9.5, BLUE, bold=True, caps_spacing=True)

def official_entry(o):
    tbl = doc.add_table(rows=1, cols=2)
    _no_borders(tbl); tbl.autofit = False
    tbl.columns[0].width = Inches(0.80); tbl.columns[1].width = Inches(5.95)
    left, right = tbl.rows[0].cells
    left.width = Inches(0.80); right.width = Inches(5.95)
    lp = tight(left.paragraphs[0])
    if o.get('portrait_file'):
        lp.add_run().add_picture(ROOT + o['portrait_file'], width=Inches(0.62))
    p = tight(right.paragraphs[0], 0, 0)
    run(p, o['name'], 10.5, NAVY, bold=True)
    if o.get('is_acting'): run(p, '  · acting', 9, ORANGE, bold=True, italic=True)
    run(p, f",  {o['title']}", 9.5, GRAY, italic=True)
    meta = o['ministry'] + (f"  ·  Assumed office: {o['assumed_office']}" if o.get('assumed_office') else '')
    pm = tight(right.add_paragraph(), 0, 1); run(pm, meta, 8, MUTED)
    pb = tight(right.add_paragraph(), 0, 0); run(pb, o.get('bio_display') or o.get('bio', ''), 9.5, INK)
    spacer = tight(doc.add_paragraph(), 0, 2); run(spacer, ' ', 2)   # minimal gap; keeps adjacent tables separate

def section(offs, div_note, label):
    shown = [o for o in offs if o.get('display') in ('full', 'half')]
    notes = [o for o in offs if o.get('display') == 'note']
    subhead(label)
    for o in shown:
        official_entry(o)
    line = div_note or ('Also: ' + '; '.join(f"{o['name']} ({o['title']})" for o in notes) if notes else '')
    if line:
        p = tight(doc.add_paragraph(), 0, 2); run(p, line, 8, GRAY, italic=True)

def key_figures(items):
    kf = tight(doc.add_paragraph(), 2, 1); kf.paragraph_format.keep_with_next = True
    run(kf, 'KEY FIGURES', 9, NAVY, bold=True, caps_spacing=True)
    rows = (len(items) + 1) // 2
    tbl = doc.add_table(rows=rows, cols=2)
    _no_borders(tbl); tbl.autofit = False
    for col in tbl.columns: col.width = Inches(3.40)
    for i, item in enumerate(items):
        cell = tbl.rows[i // 2].cells[i % 2]; cell.width = Inches(3.40)
        p = tight(cell.paragraphs[0], 0, 1)
        run(p, '▪ ', 9, GOLD, bold=True); run(p, item, 9.5, INK)
    sp = tight(doc.add_paragraph(), 0, 1); run(sp, ' ', 2)

def figures_block(c, is_us=False, au=False):
    gy, ty, oy = c['gdp_year'], c['us_trade_year'], c['overall_balance_year']
    ob = c['overall_balance_usd_b']
    basis = 'goods' if 'goods-only' in c['overall_balance_basis'] else 'goods-and-services'
    intro = tight(doc.add_paragraph(), 3, 2)
    name = c['country']
    if is_us and c.get('us_world_exports_b') is not None:
        txt = (f"The United States had a nominal GDP of {money(c['gdp_usd_b'])} in {gy} and recorded an overall goods "
               f"trade deficit of {money(ob)} in {oy}. Because the United States is the reference country for the "
               f"bilateral figures in this list, world totals are shown instead.")
        items = [f"Nominal GDP: {money_b(c['gdp_usd_b'], gy)}",
                 f"Goods balance — world: {money_b(ob, oy)} (deficit)",
                 f"Goods exports — world: {money_b(c['us_world_exports_b'], ty)}",
                 f"Goods imports — world: {money_b(c['us_world_imports_b'], ty)}"]
    elif au:
        txt = (f"African Union member states had an aggregate nominal GDP of approximately {money(c['gdp_usd_b'])} in {gy}. "
               f"No clean continental goods-trade balance is published. U.S.–Africa goods trade (a proxy for the AU) is "
               f"shown for {ty}.")
        items = [f"Aggregate nominal GDP: {money_b(c['gdp_usd_b'], gy)}",
                 "Continental goods balance: N/A",
                 f"U.S.–Africa goods exports: {money_b(c['us_exports_usd_b'], ty)}",
                 f"U.S.–Africa goods imports: {money_b(c['us_imports_usd_b'], ty)}",
                 f"U.S.–Africa goods {balance_word(c['us_bilateral_usd_b'])}: {money_b(c['us_bilateral_usd_b'], ty)}"]
    else:
        txt = (f"{name} had a nominal GDP of {money(c['gdp_usd_b'])} in {gy} and recorded an overall {basis} trade "
               f"{balance_word(ob)} of {money(ob)} in {oy}. The United States exported {money(c['us_exports_usd_b'])} in "
               f"goods to {name} and imported {money(c['us_imports_usd_b'])}, for a U.S. goods trade "
               f"{balance_word(c['us_bilateral_usd_b'])} of {money(c['us_bilateral_usd_b'])} in {ty}.")
        items = [f"Nominal GDP: {money_b(c['gdp_usd_b'], gy)}",
                 f"Overall {basis} {balance_word(ob)}: {money_b(ob, oy)}",
                 f"U.S. goods exports: {money_b(c['us_exports_usd_b'], ty)}",
                 f"U.S. goods imports: {money_b(c['us_imports_usd_b'], ty)}",
                 f"U.S. bilateral {balance_word(c['us_bilateral_usd_b'])}: {money_b(c['us_bilateral_usd_b'], ty)}"]
    run(intro, txt, 9.5, INK)
    key_figures(items)

# ============================ BODY ============================
for c in CO:
    is_us = c['country'] == 'United States'
    country_heading(c['country'], flag=c.get('flag_file'))
    figures_block(c, is_us=is_us)
    n_tr = len([o for o in c['trade_ministers'] if o.get('display') in ('full', 'half')])
    n_dg = len([o for o in c['digital_ministers'] if o.get('display') in ('full', 'half')])
    section(c['trade_ministers'], c.get('div_note_trade'), 'Trade Minister' if n_tr == 1 else 'Trade Ministers')
    section(c['digital_ministers'], c.get('div_note_digital'), 'Digital Minister' if n_dg == 1 else 'Digital Ministers')

p = doc.add_paragraph(); tight(p, 0, 2); page_break_before(p)
bottom_border(p, '0A314D', size=8, space=2)
run(p, 'G20 BLOC MEMBERS — ANNEX', 12, NAVY, bold=True, caps_spacing=True)
pa = tight(doc.add_paragraph(), 0, 4)
run(pa, 'The European Union and African Union hold G20 seats as supranational blocs and are profiled here '
        'alongside the 19 sovereign members.', 9, GRAY, italic=True)
for b in BL:
    au = b['country'] == 'African Union'
    country_heading(b['country'], tag='G20 BLOC MEMBER', flag=b.get('flag_file'), new_page=(b is not BL[0]))
    pf = tight(doc.add_paragraph(), 0, 2)
    run(pf, '   ·   '.join(f"{f['label']}: {f['value']}" for f in b.get('facts', [])), 8.5, GRAY)
    figures_block(b, au=au)
    section(b['trade_ministers'], b.get('div_note_trade'), 'Trade Lead')
    section(b['digital_ministers'], b.get('div_note_digital'), 'Digital / Technology Lead')

# ============================ FOOTER ============================
foot = sec.footer.paragraphs[0]
foot.paragraph_format.tab_stops.add_tab_stop(Inches(6.8), WD_TAB_ALIGNMENT.RIGHT)
run(foot, 'U.S. Department of Commerce  |  International Trade Administration', 8, BLUE, bold=True)
r = foot.add_run('\ttrade.gov  ·  '); r.font.name = FONT; r.font.size = Pt(8)
r.font.color.rgb = BLUE; r.bold = True
fld = OxmlElement('w:fldSimple'); fld.set(qn('w:instr'), 'PAGE')
rr = OxmlElement('w:r'); rPr = OxmlElement('w:rPr')
sz = OxmlElement('w:sz'); sz.set(qn('w:val'), '16'); rPr.append(sz)
col = OxmlElement('w:color'); col.set(qn('w:val'), '757D87'); rPr.append(col)
rr.append(rPr); fld.append(rr); foot._p.append(fld)

doc.save(ROOT + 'output/g20-ministers-trade-list.docx')
print('Wrote output/g20-ministers-trade-list.docx —', len(CO), 'countries +', len(BL), 'blocs (one page each)')
