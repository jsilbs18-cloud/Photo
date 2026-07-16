#!/usr/bin/env python3
import json, re, sys
from openpyxl import load_workbook
from pptx import Presentation
from pptx.util import Emu
import fitz
from PIL import ImageFont

ROOT='/home/user/Photo/'
D=json.load(open(ROOT+'data/g20.json')); CO=D['countries']
RES={x['country'].split(' (')[0]:x for x in json.load(open(ROOT+'data/_research.json'))['countries'] if x}
US_CONTAM=('Lutnick','Greer')

PASS=[]; FAIL=[]
def ok(c,m): (PASS if c else FAIL).append(m);
def check(c,m):
    (PASS if c else FAIL).append(('PASS' if c else 'FAIL')+': '+m)

def money(v): return 'N/A' if v is None else f"${v:,.1f}B"
def absmoney(v): return None if v is None else f"${abs(v):,.1f}B"

# ---------- 1. scope ----------
names=[c['country'] for c in CO]
check(len(CO)==19, f"exactly 19 countries (got {len(CO)})")
check('South Africa' not in names, "South Africa absent")
check('Poland' in names, "Poland present")
check(names==sorted(names), "countries alphabetical")

# ---------- 2. bilateral arithmetic ties to exports-imports ----------
for c in CO:
    if c['country']=='United States':
        check(c['us_exports_usd_b'] is None and c['us_bilateral_usd_b'] is None, "US bilateral columns N/A")
    else:
        e,i,b=c['us_exports_usd_b'],c['us_imports_usd_b'],c['us_bilateral_usd_b']
        check(round(e-i,1)==b, f"{c['country']}: bilateral {b} == {e}-{i}")

# ---------- 3. every dollar figure carries a year ----------
for c in CO:
    check(bool(c['gdp_year']) and bool(c['us_trade_year']) and bool(c['overall_balance_year']),
          f"{c['country']}: all figures carry a reference year")

# ---------- 4. WORKBOOK matches JSON ----------
wb=load_workbook(ROOT+'output/g20-ministers-trade-reference.xlsx'); ws=wb.active
# map header->col
hdr={ws.cell(1,j).value.replace('\n',' '):j for j in range(1,17)}
def wbnum(r,key):
    v=ws.cell(r,hdr[key]).value
    return v
row_by_country={ws.cell(r,1).value:r for r in range(2,21)}
for c in CO:
    r=row_by_country[c['country']]
    gcol=[k for k in hdr if k.startswith('GDP (')][0]
    check(abs((wbnum(r,gcol) or 0)-c['gdp_usd_b'])<0.05, f"xlsx GDP {c['country']}")
    # bilateral
    bcol=[k for k in hdr if k.startswith('U.S. bilateral')][0]
    bv=wbnum(r,bcol)
    if c['us_bilateral_usd_b'] is None:
        check(bv=='N/A', f"xlsx US bilateral N/A ({c['country']})")
    else:
        check(abs(bv-c['us_bilateral_usd_b'])<0.05, f"xlsx bilateral {c['country']}")
    # confidence flag
    fcol=[k for k in hdr if k.startswith('Confidence')][0]
    check(ws.cell(r,hdr[fcol]).value==c['row_confidence'], f"xlsx flag {c['country']}")
    # primary digital & trade minister name present
    dprim=next(o for o in c['digital_ministers'] if o['is_primary'])
    tprim=next(o for o in c['trade_ministers'] if o['is_primary'])
    dcol=[k for k in hdr if k.startswith('Digital / technology')][0]
    tcol=[k for k in hdr if k.startswith('Trade / commerce')][0]
    check(dprim['name'] in str(ws.cell(r,hdr[dcol]).value), f"xlsx digital name {c['country']}")
    check(tprim['name'] in str(ws.cell(r,hdr[tcol]).value), f"xlsx trade name {c['country']}")
check(ws.freeze_panes=='A2', "xlsx header row frozen")
# conditional formatting present on bilateral column
check(len(ws.conditional_formatting._cf_rules)>=1, "xlsx conditional formatting present")

# ---------- 5. PPTX matches JSON + flag embedding ----------
prs=Presentation(ROOT+'output/g20-ministers-trade-deck.pptx')
slides=list(prs.slides)
check(len(slides)==24, f"pptx 24 slides (got {len(slides)})")
def slide_text(sl):
    t=[]
    for sh in sl.shapes:
        if sh.has_text_frame: t.append(sh.text_frame.text)
    return "\n".join(t)
country_slides=slides[4:23]
check(len(country_slides)==19, "pptx 19 country slides")
pptx_country_names=[slide_text(s).split('\n')[0] for s in country_slides]
check('South Africa' not in ' '.join(pptx_country_names), "pptx: South Africa absent")
for idx,c in enumerate(CO):
    sl=country_slides[idx]; txt=slide_text(sl)
    check(c['country'] in txt, f"pptx slide {idx+1} is {c['country']}")
    check(money(c['gdp_usd_b']) in txt, f"pptx GDP str {c['country']}")
    if c['us_exports_usd_b'] is not None:
        check(absmoney(c['us_exports_usd_b']) in txt, f"pptx exports {c['country']}")
        check(absmoney(c['us_imports_usd_b']) in txt, f"pptx imports {c['country']}")
        check(absmoney(c['us_bilateral_usd_b']) in txt, f"pptx bilateral {c['country']}")
    else:
        check('N/A' in txt, f"pptx US N/A {c['country']}")
    dprim=next(o for o in c['digital_ministers'] if o['is_primary'])
    tprim=next(o for o in c['trade_ministers'] if o['is_primary'])
    check(dprim['name'] in txt and dprim['title'] in txt, f"pptx digital name+title {c['country']}")
    check(tprim['name'] in txt and tprim['title'] in txt, f"pptx trade name+title {c['country']}")
    # flag image embedded
    pics=[sh for sh in sl.shapes if sh.shape_type==13]
    check(len(pics)>=1, f"pptx flag embedded {c['country']}")
# US shows BOTH officials (Lutnick + Greer named)
us_txt=slide_text(country_slides[names.index('United States')])
check('Howard Lutnick' in us_txt and 'Jamieson Greer' in us_txt, "pptx US shows both Commerce Sec & USTR")

# ---------- 6. PDF matches JSON ----------
doc=fitz.open(ROOT+'output/g20-ministers-trade-deck.pdf')
check(doc.page_count==24, f"pdf 24 pages (got {doc.page_count})")
for i in range(doc.page_count):
    w,h=doc[i].rect.width/72, doc[i].rect.height/72
    if not (abs(w-13.333)<0.05 and abs(h-7.5)<0.05):
        check(False,f"pdf page {i+1} size {w:.2f}x{h:.2f}")
pdf_country=[doc[p].get_text() for p in range(4,23)]
for idx,c in enumerate(CO):
    txt=pdf_country[idx]
    check(c['country'] in txt, f"pdf page is {c['country']}")
    check(money(c['gdp_usd_b']) in txt, f"pdf GDP {c['country']}")
    if c['us_exports_usd_b'] is not None:
        check(absmoney(c['us_bilateral_usd_b']) in txt, f"pdf bilateral {c['country']}")
    dprim=next(o for o in c['digital_ministers'] if o['is_primary'])
    check(dprim['name'] in txt, f"pdf digital name {c['country']}")
us_pdf=pdf_country[names.index('United States')]
check('Howard Lutnick' in us_pdf and 'Jamieson Greer' in us_pdf, "pdf US shows both officials")

# ---------- 7. diacritics survive into all formats ----------
DIA=['Türkiye','Darío Leandro Genua','José Antonio Peña Merino','Katherina Reiche',
     'Maroš Šefčovič','Mehmet Fatih Kacır','Uraloğlu','Andrzej Domański',
     'Michał Baranowski','Krzysztof Gawkowski','Márcio Fernando Elias Rosa']
allpptx="\n".join(slide_text(s) for s in slides)
allpdf="\n".join(doc[p].get_text() for p in range(doc.page_count))
for name in DIA:
    check(name in allpptx, f"pptx diacritics: {name}")
    check(name in allpdf, f"pdf diacritics: {name}")

# ---------- 8. minister name/title/ministry char-for-char vs research ----------
def clean_trade(country,offs):
    if country.startswith('United States'): return offs
    return [o for o in offs if not any(u in o.get('name','') for u in US_CONTAM)]
for c in CO:
    rc=RES[c['country']]
    rset={(o['name'],o.get('title',''),o.get('ministry','')) for o in rc['digital_officials']+clean_trade(c['country'],rc['trade_officials'])}
    for grp in ('digital_ministers','trade_ministers'):
        for o in c[grp]:
            key=(o['name'],o['title'],o['ministry'])
            check(key in rset, f"research match: {c['country']} / {o['name']}")

# ---------- 9. flags valid ----------
from PIL import Image
for c in CO:
    try:
        im=Image.open(ROOT+c['flag_file']); im.load()
        check(im.height>0 and im.width>0, f"flag valid {c['country']}")
    except Exception as e:
        check(False, f"flag load {c['country']}: {e}")

# ---------- 10. PPTX text-fit (overflow) measurement ----------
FP={'reg':'/usr/share/fonts/truetype/liberation/LiberationSans-Regular.ttf',
    'bold':'/usr/share/fonts/truetype/liberation/LiberationSans-Bold.ttf',
    'ital':'/usr/share/fonts/truetype/liberation/LiberationSans-Italic.ttf'}
_fc={}
def font(style,pt):
    px=max(6,round(pt*96/72)); k=(style,px)
    if k not in _fc: _fc[k]=ImageFont.truetype(FP[style],px)
    return _fc[k]
def nlines(text,pt,width_in,style='reg'):
    if not text: return 0
    f=font(style,pt); maxw=width_in*96; lines=1; cur=0.0
    space=f.getlength(' ')
    for word in text.split():
        wl=f.getlength(word)
        if cur==0: cur=wl
        elif cur+space+wl<=maxw: cur+=space+wl
        else: lines+=1; cur=wl
    return lines
RW=8.23
worst=[]
for c in CO:
    for grp,bh in (('digital_ministers',2.82),('trade_ministers',2.82)):
        o=next(x for x in c[grp] if x['is_primary'])
        frame_h=bh-1.44
        bio_lines=nlines(o['bio'],11,RW,'reg')
        bio_h=bio_lines*11*1.03/72
        note=''
        extra=[x for x in c[grp] if not x['is_primary']]
        bespoke=c.get('div_note_'+('digital' if grp=='digital_ministers' else 'trade'))
        if bespoke: note=bespoke
        elif extra: note='also '+'; '.join(f"{x['name']} ({x['title']})" for x in extra)
        note_full=("Division of responsibility:  "+note) if note else ''
        note_h=0
        if note_full:
            note_h=6/72 + nlines(note_full,9.5,RW,'ital')*9.5*1.2/72
        total=bio_h+note_h
        # title & ministry fit
        tl=nlines(o['title'],11,RW,'ital'); ml=nlines(o['ministry']+'    Assumed office: '+o.get('assumed_office',''),10,RW,'reg')
        fits = total<=frame_h+0.02 and tl<=2 and ml<=2
        worst.append((total/frame_h, c['country'], grp, bio_lines, round(total,2), round(frame_h,2), tl, ml))
        check(fits, f"pptx fit {c['country']}/{grp[:3]} bio+note={total:.2f}in<= {frame_h:.2f} (title {tl}L, min {ml}L)")

# ---------- 11. GAPS.md non-empty ----------
import os
g=ROOT+'output/GAPS.md'
check(os.path.exists(g) and os.path.getsize(g)>500, "GAPS.md exists & non-empty")

# ---------- report ----------
print(f"\n{'='*60}\nQA RESULTS: {len(PASS)} passed, {len(FAIL)} failed\n{'='*60}")
if FAIL:
    print("\nFAILURES:")
    for f in FAIL: print("  ✗", f if isinstance(f,str) else f)
else:
    print("\nALL CHECKS PASSED ✓")
worst.sort(reverse=True)
print("\nTightest text boxes (ratio of content to frame height):")
for r in worst[:6]:
    print(f"  {r[1]:<16}{r[2][:3]}  ratio={r[0]:.2f}  bio={r[3]}L  content={r[4]}in/{r[5]}in  title={r[6]}L min={r[7]}L")
sys.exit(1 if FAIL else 0)
